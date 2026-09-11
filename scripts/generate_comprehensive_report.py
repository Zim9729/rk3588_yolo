import base64
import html
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import cv2
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EVAL_XLSX = ROOT / "outputs/final_summary/dataset_eval/precision_loss_eval.xlsx"
BENCH_XLSX = ROOT / "outputs/final_summary/single_image/precision_benchmark.xlsx"
BENCH_DIR = ROOT / "outputs/final_summary/single_image"
REPORT_DIR = ROOT / "reports"
MD_PATH = REPORT_DIR / "RKNN_COMPREHENSIVE_EVALUATION.md"
HTML_PATH = REPORT_DIR / "RKNN_COMPREHENSIVE_EVALUATION.html"

MODEL_PATHS = {
    "ARM-A-FP16": "models_arm/3C_stage1_pathA_fp16.rknn",
    "ARM-A-INT8": "models_arm/3C_stage1_pathA_int8.rknn",
    "ARM-B-FP16": "models_arm/3C_stage1_fp16.rknn",
    "ARM-B-INT8": "models_arm/3C_stage1_int8.rknn",
    "ARM-B-INT8-layer": "models_arm/3C_stage1_int8_layer.rknn",
    "ARM-B-Hybrid": "models_arm/3C_stage1_hybrid.rknn",
    "WSL2-A-FP16": "models_wsl2/3C_SDG_stage1_1_best_v2_pathA_fp16.rknn",
    "WSL2-A-INT8": "models_wsl2/3C_SDG_stage1_1_best_v2_pathA_int8.rknn",
    "WSL2-B-FP16": "models_wsl2/3C_SDG_stage1_1_best_v2_pathB_fp16.rknn",
    "WSL2-B-INT8": "models_wsl2/3C_SDG_stage1_1_best_v2_pathB_int8.rknn",
    "WSL2-B-INT8-layer": "models_wsl2/3C_SDG_stage1_1_best_v2_pathB_int8_layer.rknn",
    "WSL2-B-Hybrid": "models_wsl2/3C_SDG_stage1_1_best_v2_pathB_hybrid.rknn",
    "WSL2-B-INT8-mmse": "models_wsl2/3C_SDG_stage1_1_best_v2_pathB_int8_mmse.rknn",
    "WSL2-B-INT8-kl": "models_wsl2/3C_SDG_stage1_1_best_v2_pathB_int8_kl.rknn",
}


def read_table(path, sheet, header_value):
    ws = load_workbook(path, data_only=True, read_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    index = next(i for i, row in enumerate(rows) if row and row[0] == header_value)
    headers = rows[index]
    return [dict(zip(headers, row)) for row in rows[index + 1 :] if row and row[0] is not None]


def classify(name):
    platform = "板端转换" if name.startswith("ARM-") else "WSL2 转换"
    path = "Path A" if "-A-" in name else "Path B"
    if name.endswith("FP16"):
        precision = "FP16"
        algorithm = "不量化"
    elif name.endswith("INT8-layer"):
        precision = "INT8"
        algorithm = "normal / per-tensor"
    elif name.endswith("Hybrid"):
        precision = "Hybrid"
        algorithm = "normal / auto-hybrid"
    elif name.endswith("INT8-mmse"):
        precision = "INT8"
        algorithm = "mmse / per-channel"
    elif name.endswith("INT8-kl"):
        precision = "INT8"
        algorithm = "kl_divergence / per-channel"
    else:
        precision = "INT8"
        algorithm = "normal / per-channel"
    return platform, path, precision, algorithm


def image_data_uri(path, max_width=720, quality=76):
    image = cv2.imread(str(path))
    if image is None:
        return ""
    height, width = image.shape[:2]
    if width > max_width:
        scale = max_width / width
        image = cv2.resize(image, (max_width, round(height * scale)))
    ok, data = cv2.imencode(".jpg", image, [cv2.IMWRITE_JPEG_QUALITY, quality])
    if not ok:
        return ""
    return "data:image/jpeg;base64," + base64.b64encode(data.tobytes()).decode("ascii")


def fmt(value, digits=4):
    return f"{float(value):.{digits}f}"


def model_records():
    eval_rows = {row["精度"]: row for row in read_table(EVAL_XLSX, "Summary", "精度")}
    bench_rows = {row["精度"]: row for row in read_table(BENCH_XLSX, "Summary", "精度")}
    accuracy_rows = {row["精度"]: row for row in read_table(BENCH_XLSX, "Accuracy vs FP16", "精度")}
    detection_rows = read_table(BENCH_XLSX, "Detections", "精度")
    detections = defaultdict(list)
    for row in detection_rows:
        detections[row["精度"]].append(row)
    records = []
    for name, relative_path in MODEL_PATHS.items():
        platform, path, precision, algorithm = classify(name)
        model_path = ROOT / relative_path
        evaluation = eval_rows[name]
        benchmark = bench_rows[name]
        accuracy = accuracy_rows[name]
        records.append({
            "name": name,
            "platform": platform,
            "path": path,
            "precision_type": precision,
            "algorithm": algorithm,
            "model_path": relative_path,
            "size_mb": model_path.stat().st_size / (1024 * 1024),
            "tp": evaluation["TP"],
            "fp": evaluation["FP"],
            "fn": evaluation["FN"],
            "precision": evaluation["Precision"],
            "recall": evaluation["Recall"],
            "f1": evaluation["F1"],
            "avg_iou": evaluation["平均IoU"],
            "avg_score": evaluation["平均score"],
            "dataset_avg_ms": evaluation["推理avg(ms)"],
            "dataset_p50_ms": evaluation["推理p50(ms)"],
            "single_count": benchmark["检测数"],
            "single_avg_ms": benchmark["推理 avg(ms)"],
            "single_p50_ms": benchmark["推理 p50(ms)"],
            "single_std_ms": benchmark["推理 std(ms)"],
            "single_total_ms": benchmark["总耗时 avg(ms)"],
            "single_iou_vs_fp16": accuracy["平均IoU"],
            "single_score_diff": accuracy["平均分数差"],
            "single_box_diff": accuracy["平均框坐标差(px)"],
            "single_unmatched": accuracy["未匹配(其他)"],
            "detections": detections[name],
        })
    return records


def markdown_report(records):
    fp16 = next(r for r in records if r["name"] == "WSL2-B-FP16")
    best = next(r for r in records if r["name"] == "WSL2-B-INT8")
    speedup = fp16["single_avg_ms"] / best["single_avg_ms"]
    size_reduction = 1 - best["size_mb"] / fp16["size_mb"]
    rows = []
    for r in records:
        single = "1（正确）" if r["single_count"] == 1 else f"{r['single_count']}（重复框）"
        rows.append(
            f"| {r['name']} | {r['platform']} | {r['path']} | {r['algorithm']} | {r['size_mb']:.2f} | "
            f"{r['tp']} / {r['fp']} / {r['fn']} | {r['precision']:.4f} | {r['recall']:.4f} | {r['f1']:.4f} | "
            f"{r['avg_iou']:.4f} | {r['avg_score']:.4f} | {r['dataset_avg_ms']:.2f} | {r['dataset_p50_ms']:.2f} | "
            f"{single} | {r['single_avg_ms']:.2f} | {r['single_p50_ms']:.2f} |"
        )
    arm_normal = next(r for r in records if r["name"] == "ARM-B-INT8")
    wsl_normal = best
    mmse = next(r for r in records if r["name"] == "WSL2-B-INT8-mmse")
    kl = next(r for r in records if r["name"] == "WSL2-B-INT8-kl")
    wsl_layer = next(r for r in records if r["name"] == "WSL2-B-INT8-layer")
    arm_layer = next(r for r in records if r["name"] == "ARM-B-INT8-layer")
    gallery = "\n".join(
        f"- `{r['name']}`：`../outputs/final_summary/single_image/{r['name']}_result.jpg`"
        for r in records
    )
    return f"""# RK3588 YOLO11 RKNN 全量评测总结

> 模型：`3C_SDG_stage1_1_best_v2.pt`  
> 类别：`Pantograph_Area`（1 类）  
> 评测对象：板端转换与 WSL2 转换的 Path A / Path B 全部有效 RKNN 模型  
> 报告生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

## 1. 最终结论

### 最佳综合模型

**推荐：`WSL2-B-INT8`**  
模型文件：`{best['model_path']}`

选择理由：

- 300 张 XML 真值集：`TP={best['tp']}`、`FP={best['fp']}`、`FN={best['fn']}`，Precision/Recall/F1 均为 `1.0000`。
- 平均 IoU 为 `{best['avg_iou']:.4f}`；相比 FP16 的 `{fp16['avg_iou']:.4f}`，绝对下降 `{fp16['avg_iou'] - best['avg_iou']:.4f}`，损失主要来自框回归，而不是分类或召回。
- `3C_Test.jpg` 只输出 **1 个框**，符合已确认的正确行为；板端转换的 normal INT8 输出 2 个重叠框。
- 预热 5 次、计时 30 次后，纯 NPU 推理平均 `{best['single_avg_ms']:.2f} ms`，相对 WSL2 Path B FP16 的 `{fp16['single_avg_ms']:.2f} ms` 加速约 `{speedup:.2f}x`。
- 模型大小 `{best['size_mb']:.2f} MiB`，相对 FP16 的 `{fp16['size_mb']:.2f} MiB` 减少约 `{size_reduction * 100:.1f}%`。
- Path B 保留 ONNX 中间产物，并允许选择量化算法和粒度，复现和后续调优能力优于 Path A。

### 精度优先模型

若定位精度比速度更重要，选择 **`WSL2-B-FP16`** 或任一 FP16：平均 IoU `{fp16['avg_iou']:.4f}`，300 张图 0 误检、0 漏检，单图结果正确；代价是推理约为 INT8 的 `{speedup:.2f}x`，模型也更大。

### 不推荐模型

- `WSL2-B-INT8-mmse`：平均 IoU `{mmse['avg_iou']:.4f}`，出现 `{mmse['fp']}` 个 FP；本数据集上明显劣于 normal。
- `WSL2-B-INT8-kl`：平均 IoU `{kl['avg_iou']:.4f}`，且 `3C_Test.jpg` 输出 2 个框。
- `WSL2-B-INT8-layer`：平均 IoU `{wsl_layer['avg_iou']:.4f}` 为 INT8 中最高，但 300 图出现 `{wsl_layer['fp']}` 个 FP，综合稳定性不如 normal per-channel。
- 板端量化的 INT8/Hybrid：全量集指标可用，但在目标单图上均产生重复框。

## 2. 统一测试条件

| 项目 | 配置 |
|---|---|
| RK3588 Runtime | `librknnrt 1.6.0`，driver `0.9.8` |
| RKNN Toolkit / 模型版本 | `2.3.2` |
| 输入尺寸 | `640 × 640` |
| 类别 | `Pantograph_Area` |
| 置信度阈值 | `0.25` |
| NMS IoU 阈值 | `0.45` |
| 全量评测 | `data/calibration` 中 300 张图，真值来自 `data/xml` |
| TP 匹配阈值 | IoU ≥ `0.5` |
| 单图基准 | `3C_Test.jpg`，warmup=5，runs=30 |
| 正确单图行为 | 1 个检测框（由人工确认） |

说明：单图表中的延迟是预热后的**纯 NPU 推理时间**。300 图评测脚本记录的 `推理avg` 实际覆盖预处理、NPU 推理和后处理，因此不作为主要速度排名依据。测试期间存在 `RKNN model 2.3.2` 与板端 runtime `1.6.0` 的版本不匹配警告，生产部署应优先升级 Runtime/驱动后复测。

## 3. 全部模型综合结果

| 模型 | 转换环境 | 路径 | 量化方式 | 大小 MiB | TP/FP/FN | P | R | F1 | 平均IoU | 平均score | 300图端到端avg ms | 300图p50 ms | 单图检测数 | 单图NPU avg ms | p50 ms |
|---|---|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|
{chr(10).join(rows)}

## 4. 300 张 XML 真值集分析

### 4.1 FP16

所有 FP16 模型的结果一致：`TP=300`、`FP=0`、`FN=0`、平均 IoU `0.9276`、平均 score `0.8560`。这说明 Path A 和 Path B 在浮点模型上没有可观测精度差异。

### 4.2 normal INT8

- 板端 Path A：IoU `0.8628`；板端 Path B：IoU `{arm_normal['avg_iou']:.4f}`。
- WSL2 Path A：IoU `0.8624`；WSL2 Path B：IoU `{wsl_normal['avg_iou']:.4f}`。
- 四个 normal INT8 在 300 图上均为 0 FP、0 FN，IoU 差异最大仅约 `0.0010`，整体精度近似等价。
- 量化损失集中在定位：平均 IoU 相比 FP16 下降约 `0.0643–0.0653`，平均 score 只下降约 `0.009`。

### 4.3 layer、Hybrid、MMSE 与 KL

- 板端 `INT8-layer`：IoU `{arm_layer['avg_iou']:.4f}`，300 图无 FP，但目标单图有重复框。
- WSL2 `INT8-layer`：IoU `{wsl_layer['avg_iou']:.4f}`，略高于 normal，但出现 1 个 FP。
- WSL2 Hybrid：IoU `0.8619`，速度也没有优于 normal，自动混合精度没有带来收益。
- MMSE：IoU `{mmse['avg_iou']:.4f}`，较 normal 再下降 `{best['avg_iou'] - mmse['avg_iou']:.4f}`，平均 score 也降至 `{mmse['avg_score']:.4f}`。MMSE 并不保证检测模型更准，本模型检测头存在权重离群值，且其最小均方误差目标不等价于最大化最终框 IoU。
- KL：IoU `{kl['avg_iou']:.4f}`，低于 normal，目标单图也产生重复框。

## 5. `3C_Test.jpg` 单图分析

### 5.1 检测数

- 板端转换：FP16 为 1 个框；Path A INT8、Path B INT8、INT8-layer、Hybrid 均为 **2 个框**。
- WSL2 转换：FP16、Path A INT8、Path B normal INT8、INT8-layer、Hybrid、MMSE 均为 **1 个框**；只有 KL 为 **2 个框**。
- 板端 normal INT8 的两个框 IoU 约 `0.4386`，略低于 NMS 阈值 `0.45`，因此没有互相抑制。

### 5.2 主要框坐标

| 模型组 | score | box `(x1, y1, x2, y2)` | 与 FP16 框 IoU |
|---|---:|---|---:|
| FP16 | 0.8110 | `(654.0, 1582.5, 1886.0, 1833.5)` | 1.0000 |
| normal INT8 主框 | 0.8184 | `(637.0, 1612.2, 1911.0, 1881.2)` | 0.7189 |
| 板端 normal INT8 重复框 | 0.7409 | `(679.5, 1513.1, 1868.5, 1782.1)` | 约 0.6058 |
| WSL2 MMSE 主框 | 0.8153 | `(715.4, 1622.7, 1883.7, 1885.2)` | 0.6661 |
| WSL2 KL 最佳匹配框 | 0.8007 | `(678.5, 1553.0, 1950.8, 1821.6)` | 0.7972 |

KL 的最佳匹配框与 FP16 更接近，但它仍输出第二个高分框，因此不能只看单个最佳框 IoU。MMSE 虽然只输出一个框，但定位误差更大，这与其 300 图平均 IoU 最低相一致。

### 5.3 单图输出图位置

{gallery}

HTML 版报告已内嵌全部输出图，可直接离线展示。

## 6. Path A 与 Path B 对比

| 维度 | Path A | Path B |
|---|---|---|
| 工作流 | Ultralytics `yolo export format=rknn` 一步完成 | `.pt → ONNX → RKNN` 分步完成 |
| FP16 精度 | 与 Path B 相同 | 与 Path A 相同 |
| normal INT8 精度 | 与 Path B 基本相同 | 与 Path A 基本相同 |
| 中间 ONNX | 一般不作为稳定产物保留 | 明确保留，便于复现和诊断 |
| 算法/粒度控制 | 较少 | normal/mmse/KL、channel/layer、Hybrid |
| 推荐用途 | 快速转换 | 正式部署、实验对比和后续优化 |

## 7. 板端转换与 WSL2 转换对比

- 300 图整体指标几乎相同：板端 Path B normal INT8 IoU `{arm_normal['avg_iou']:.4f}`，WSL2 为 `{wsl_normal['avg_iou']:.4f}`。
- 目标单图行为不同：板端 normal INT8 为 2 框，WSL2 normal INT8 为 1 框。
- 观测到模型编译器构建不同：板端模型为 `2.3.2 (@2025-04-03...)`，WSL2 模型为 `2.3.2 (e045de294f@2025-04-07...)`。但不能仅据此断定差异完全由编译器造成；校准图片遍历顺序、ONNX 产物、Toolkit wheel 构建和转换环境也可能影响量化参数。
- 基于本次实际结果，**WSL2 生成的 normal per-channel INT8 更适合部署**。

## 8. 最终选择矩阵

| 目标 | 推荐模型 | 理由 |
|---|---|---|
| 综合部署（首选） | `WSL2-B-INT8` | 300 图 0 FP/0 FN；单图 1 框；速度快；体积小；流程可复现 |
| 定位精度最高 | `WSL2-B-FP16` | IoU 0.9276，0 FP/0 FN，单图正确 |
| 只追求单图最低延迟 | 不建议仅按最低数字选型 | 板端 INT8 虽快，但目标单图重复框；应先满足正确性 |
| 快速简单转换 | `WSL2-A-INT8` | 与 Path B normal 基本等价，且单图正确 |
| 继续量化研究 | 以 `WSL2-B-INT8` 为基线 | 当前 MMSE/KL/Hybrid 均未改善综合指标 |

## 9. 原始产物

- 统一 300 图 Excel：`../outputs/final_summary/dataset_eval/precision_loss_eval.xlsx`
- 统一 300 图原始 HTML：`../outputs/final_summary/dataset_eval/precision_loss_eval.html`
- 统一单图 Excel：`../outputs/final_summary/single_image/precision_benchmark.xlsx`
- 统一单图原始 HTML：`../outputs/final_summary/single_image/precision_benchmark.html`
- 本综合报告 HTML：`RKNN_COMPREHENSIVE_EVALUATION.html`
"""


def bar_chart(records, key, value_format, max_value=None, inverse=False):
    values = [float(r[key]) for r in records]
    upper = max_value or max(values)
    lines = []
    for r, value in zip(records, values):
        width = min(100, value / upper * 100 if upper else 0)
        if inverse:
            color = "#ef6c67" if value > 1 else "#31b77a"
        elif r["name"] == "WSL2-B-INT8":
            color = "#18a77b"
        elif r["precision_type"] == "FP16":
            color = "#6a7de1"
        else:
            color = "#4da3d9"
        lines.append(
            f'<div class="bar-row"><span class="bar-label">{html.escape(r["name"])}</span>'
            f'<div class="bar-track"><span class="bar-fill" style="width:{width:.1f}%;background:{color}"></span></div>'
            f'<strong>{value_format(value)}</strong></div>'
        )
    return "".join(lines)


def html_report(records):
    fp16 = next(r for r in records if r["name"] == "WSL2-B-FP16")
    best = next(r for r in records if r["name"] == "WSL2-B-INT8")
    speedup = fp16["single_avg_ms"] / best["single_avg_ms"]
    size_reduction = 1 - best["size_mb"] / fp16["size_mb"]
    table_rows = []
    for r in records:
        state = "ok" if r["fp"] == 0 and r["fn"] == 0 and r["single_count"] == 1 else "warn"
        badge = "推荐" if r["name"] == "WSL2-B-INT8" else ("正确" if state == "ok" else "需注意")
        table_rows.append(f"""
<tr class="{state} {'best-row' if r['name'] == 'WSL2-B-INT8' else ''}">
<td><strong>{html.escape(r['name'])}</strong><span class="tag {state}">{badge}</span></td>
<td>{r['platform']}</td><td>{r['path']}</td><td>{r['algorithm']}</td><td>{r['size_mb']:.2f}</td>
<td>{r['tp']}/{r['fp']}/{r['fn']}</td><td>{r['precision']:.4f}</td><td>{r['recall']:.4f}</td><td>{r['f1']:.4f}</td>
<td>{r['avg_iou']:.4f}</td><td>{r['avg_score']:.4f}</td><td>{r['dataset_avg_ms']:.2f}</td><td>{r['dataset_p50_ms']:.2f}</td><td>{r['single_count']}</td>
<td>{r['single_avg_ms']:.2f}</td><td>{r['single_p50_ms']:.2f}</td>
</tr>""")
    gallery = []
    for r in records:
        path = BENCH_DIR / f"{r['name']}_result.jpg"
        uri = image_data_uri(path)
        det_lines = "".join(
            f"<li>score {d['score']:.4f} · ({d['x1']:.1f}, {d['y1']:.1f}, {d['x2']:.1f}, {d['y2']:.1f})</li>"
            for d in r["detections"]
        )
        status = "正确：1 框" if r["single_count"] == 1 else f"重复检测：{r['single_count']} 框"
        gallery.append(f"""
<article class="image-card {'bad-card' if r['single_count'] != 1 else ''}">
<div class="image-head"><strong>{html.escape(r['name'])}</strong><span>{status}</span></div>
<img loading="lazy" src="{uri}" alt="{html.escape(r['name'])} detection">
<ul>{det_lines}</ul>
</article>""")
    iou_chart = bar_chart(records, "avg_iou", lambda v: f"{v:.4f}", max_value=1)
    latency_chart = bar_chart(records, "single_avg_ms", lambda v: f"{v:.2f} ms")
    size_chart = bar_chart(records, "size_mb", lambda v: f"{v:.2f} MiB")
    fp_chart = bar_chart(records, "fp", lambda v: f"{int(v)} FP", max_value=1, inverse=True)
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>RK3588 YOLO11 RKNN 全量评测</title>
<style>
:root{{--bg:#07111f;--panel:#0e1b2d;--panel2:#13243a;--line:#263d58;--text:#e9f0f8;--muted:#97abc2;--green:#35d39a;--cyan:#55c4e8;--blue:#8095ff;--amber:#f3b95f;--red:#ef746e}}
*{{box-sizing:border-box}}html{{scroll-behavior:smooth}}body{{margin:0;background:radial-gradient(circle at 85% 0,#15304e 0,transparent 32%),var(--bg);color:var(--text);font-family:Inter,"Segoe UI","Microsoft YaHei",sans-serif;line-height:1.65}}
a{{color:#79d5f1}}.wrap{{width:min(1500px,calc(100% - 36px));margin:auto}}header{{padding:70px 0 45px;border-bottom:1px solid var(--line)}}.eyebrow{{color:var(--cyan);letter-spacing:.13em;text-transform:uppercase;font-size:13px;font-weight:700}}h1{{font-size:clamp(34px,5vw,64px);line-height:1.06;margin:12px 0 18px;max-width:1000px}}h2{{margin-top:0;font-size:27px}}h3{{font-size:19px}}.lead{{max-width:920px;color:var(--muted);font-size:18px}}nav{{position:sticky;top:0;z-index:10;background:#07111fe6;backdrop-filter:blur(14px);border-bottom:1px solid var(--line)}}nav .wrap{{display:flex;gap:22px;overflow:auto;padding:12px 0}}nav a{{white-space:nowrap;text-decoration:none;font-size:14px;color:#bcd0e4}}section{{padding:44px 0}}.grid4{{display:grid;grid-template-columns:repeat(4,1fr);gap:16px}}.metric-card,.panel,.decision{{background:linear-gradient(145deg,var(--panel2),var(--panel));border:1px solid var(--line);border-radius:16px;box-shadow:0 18px 45px #0003}}.metric-card{{padding:21px}}.metric-card small{{display:block;color:var(--muted)}}.metric-card strong{{display:block;font-size:30px;margin-top:5px}}.metric-card.best{{border-color:#299f7a;box-shadow:0 0 0 1px #299f7a44,0 18px 45px #0003}}.panel{{padding:24px;margin-bottom:22px}}.recommendation{{display:grid;grid-template-columns:1.4fr .6fr;gap:18px}}.hero-choice{{padding:28px;background:linear-gradient(135deg,#153f43,#102a38);border:1px solid #267c71;border-radius:18px}}.hero-choice code{{font-size:20px;color:#8bf2cf}}.checks{{padding:0;list-style:none}}.checks li{{padding:8px 0;border-bottom:1px solid #ffffff12}}.checks li:before{{content:"✓";color:var(--green);font-weight:bold;margin-right:10px}}.warning{{padding:18px;border-left:4px solid var(--amber);background:#3e301c88;border-radius:8px;color:#f5d7a4}}.table-wrap{{overflow:auto;border:1px solid var(--line);border-radius:14px}}table{{border-collapse:collapse;width:100%;min-width:1250px;font-size:13px;background:#0b1728}}th,td{{padding:11px 10px;border-bottom:1px solid #21364e;text-align:right;white-space:nowrap}}th{{position:sticky;top:47px;background:#172a42;color:#bad0e7;z-index:2}}th:first-child,td:first-child,th:nth-child(2),td:nth-child(2),th:nth-child(3),td:nth-child(3),th:nth-child(4),td:nth-child(4){{text-align:left}}tr:hover{{background:#12243a}}.best-row{{background:#113b35!important}}.tag{{display:block;width:max-content;font-size:10px;border-radius:20px;padding:1px 7px;margin-top:3px}}.tag.ok{{background:#173e35;color:#76dfb8}}.tag.warn{{background:#482b27;color:#ffaaa2}}.charts{{display:grid;grid-template-columns:1fr 1fr;gap:20px}}.bar-row{{display:grid;grid-template-columns:155px 1fr 82px;gap:10px;align-items:center;margin:9px 0;font-size:12px}}.bar-label{{overflow:hidden;text-overflow:ellipsis;white-space:nowrap;color:#b9cadd}}.bar-track{{height:10px;background:#07111f;border-radius:9px;overflow:hidden}}.bar-fill{{height:100%;display:block;border-radius:9px}}.bar-row strong{{font-size:12px;text-align:right}}.gallery{{display:grid;grid-template-columns:repeat(2,1fr);gap:18px}}.image-card{{background:var(--panel);border:1px solid var(--line);border-radius:14px;overflow:hidden}}.image-card.bad-card{{border-color:#9c554e}}.image-card img{{display:block;width:100%;aspect-ratio:16/9;object-fit:cover}}.image-head{{display:flex;justify-content:space-between;padding:13px 15px;background:#13243a}}.image-head span{{font-size:12px;color:var(--muted)}}.bad-card .image-head span{{color:#ff9992}}.image-card ul{{font-size:12px;color:#aebfd2;margin:10px 18px 14px;padding-left:18px}}.decision{{padding:26px;border-color:#2a8e72}}.decision-grid{{display:grid;grid-template-columns:repeat(3,1fr);gap:15px}}.decision-grid div{{background:#07111f88;border-radius:10px;padding:15px}}code{{background:#07111f;padding:2px 6px;border-radius:5px}}footer{{border-top:1px solid var(--line);padding:30px 0 60px;color:var(--muted);font-size:13px}}@media(max-width:950px){{.grid4,.charts,.recommendation,.decision-grid{{grid-template-columns:1fr 1fr}}.gallery{{grid-template-columns:1fr}}}}@media(max-width:620px){{.grid4,.charts,.recommendation,.decision-grid{{grid-template-columns:1fr}}.wrap{{width:min(100% - 22px,1500px)}}header{{padding-top:45px}}}}
</style>
</head>
<body>
<header><div class="wrap"><div class="eyebrow">RK3588 · YOLO11 · RKNN Toolkit 2.3.2</div><h1>Path A / Path B<br>板端与 WSL2 全量评测</h1><p class="lead">统一比较 14 个 RKNN 模型：300 张 XML 真值集、3C_Test.jpg 单图行为、NPU 延迟、模型体积与量化损失。所有图表均内嵌，可离线展示。</p></div></header>
<nav><div class="wrap"><a href="#conclusion">最终结论</a><a href="#method">测试条件</a><a href="#overview">完整指标</a><a href="#charts">可视化</a><a href="#single">单图结果</a><a href="#analysis">分析</a></div></nav>
<main>
<section id="conclusion"><div class="wrap"><h2>最终结论</h2><div class="recommendation"><div class="hero-choice"><span class="eyebrow">最佳综合模型</span><h2>WSL2-B-INT8</h2><code>{best['model_path']}</code><ul class="checks"><li>300 图：300 TP / 0 FP / 0 FN</li><li>3C_Test.jpg：1 个框，行为正确</li><li>平均 IoU {best['avg_iou']:.4f}，纯 NPU 推理 {best['single_avg_ms']:.2f} ms</li><li>normal per-channel，兼顾稳定性、速度和复现能力</li></ul></div><div class="decision"><h3>为什么不是 MMSE？</h3><p>MMSE 的最终优化目标是张量均方误差，不是检测框 IoU。本次 IoU 仅 {next(r for r in records if r['name']=='WSL2-B-INT8-mmse')['avg_iou']:.4f}，并出现 1 个 FP。</p><h3>为什么选择 Path B？</h3><p>精度与 Path A 基本相同，但保留 ONNX 并支持算法、量化粒度和 Hybrid 控制。</p></div></div><div class="grid4" style="margin-top:18px"><div class="metric-card best"><small>推荐模型平均 IoU</small><strong>{best['avg_iou']:.4f}</strong></div><div class="metric-card"><small>对 FP16 加速</small><strong>{speedup:.2f}×</strong></div><div class="metric-card"><small>模型体积降低</small><strong>{size_reduction*100:.1f}%</strong></div><div class="metric-card"><small>300 图 FP / FN</small><strong>0 / 0</strong></div></div></div></section>
<section id="method"><div class="wrap"><h2>统一测试条件</h2><div class="panel"><div class="grid4"><div><small>输入</small><strong>640 × 640</strong></div><div><small>阈值</small><strong>conf 0.25</strong></div><div><small>NMS</small><strong>IoU 0.45</strong></div><div><small>GT 匹配</small><strong>IoU ≥ 0.5</strong></div></div><p>全量集：<code>data/calibration</code> 300 张图，XML 来自 <code>data/xml</code>。单图：<code>3C_Test.jpg</code>，warmup=5、runs=30。单图正确行为由人工确认是 1 个检测框。</p><div class="warning"><strong>环境警告：</strong>模型由 RKNN Toolkit 2.3.2 生成，但板端 librknnrt 为 1.6.0、driver 为 0.9.8。模型可运行，但生产结论应在 Runtime/驱动升级后再次确认。全量脚本的延迟包含预处理和后处理；速度排名采用预热后的单图纯 NPU 推理时间。</div></div></div></section>
<section id="overview"><div class="wrap"><h2>14 个模型完整指标</h2><div class="table-wrap"><table><thead><tr><th>模型</th><th>环境</th><th>路径</th><th>量化方式</th><th>MiB</th><th>TP/FP/FN</th><th>P</th><th>R</th><th>F1</th><th>平均IoU</th><th>平均score</th><th>300图端到端avg</th><th>300图p50</th><th>单图框数</th><th>NPU avg</th><th>NPU p50</th></tr></thead><tbody>{''.join(table_rows)}</tbody></table></div></div></section>
<section id="charts"><div class="wrap"><h2>指标可视化</h2><div class="charts"><div class="panel"><h3>平均 IoU（越高越好）</h3>{iou_chart}</div><div class="panel"><h3>单图纯 NPU 推理（越低越好）</h3>{latency_chart}</div><div class="panel"><h3>模型体积</h3>{size_chart}</div><div class="panel"><h3>300 图误检数（越低越好）</h3>{fp_chart}</div></div></div></section>
<section id="single"><div class="wrap"><h2>3C_Test.jpg 单图检测可视化</h2><p class="lead">绿色框为各模型输出。板端 INT8/Hybrid 与 WSL2 KL 出现两个重叠检测；其余 WSL2 模型保持一个检测。</p><div class="gallery">{''.join(gallery)}</div></div></section>
<section id="analysis"><div class="wrap"><h2>综合分析</h2><div class="charts"><div class="panel"><h3>Path A vs Path B</h3><p>FP16 的平均 IoU 均为 0.9276；normal INT8 的 IoU 位于 0.8623–0.8633，路径本身不构成实质精度差异。Path B 的优势是可控和可复现，而不是本轮数值更高。</p><p>WSL2 Path B normal INT8 在单图正确性上优于板端版本，因此成为首选。</p></div><div class="panel"><h3>量化算法</h3><p>normal per-channel 的综合表现最好。per-tensor 仅提升约 0.0025 IoU，却在 WSL2 全量集产生 1 个 FP。Hybrid 没有提高 IoU，且略慢。</p><p>MMSE 和 KL 均降低定位质量；KL 还在目标单图产生重复框。</p></div><div class="panel"><h3>精度损失类型</h3><p>normal INT8 保持了 300/300 召回和 0 FP，平均 score 只下降约 0.009；主要损失来自框坐标回归，平均 IoU 比 FP16 低约 0.065。</p></div><div class="panel"><h3>环境差异</h3><p>板端模型编译器标识为 2025-04-03 构建，WSL2 为 2025-04-07 构建。单图差异可能同时来自编译器、ONNX、校准顺序或 wheel 构建，不能只归因于编译器版本。</p></div></div><div class="decision"><h2>选择矩阵</h2><div class="decision-grid"><div><strong>综合部署</strong><p>WSL2-B-INT8<br>稳定、快速、体积小</p></div><div><strong>定位精度优先</strong><p>WSL2-B-FP16<br>IoU 0.9276</p></div><div><strong>快速转换</strong><p>WSL2-A-INT8<br>一步导出且单图正确</p></div></div></div></div></section>
</main>
<footer><div class="wrap"><p>生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')} · 原始 Excel：outputs/final_summary/dataset_eval 与 outputs/final_summary/single_image</p></div></footer>
</body></html>"""


def main():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    records = model_records()
    MD_PATH.write_text(markdown_report(records), encoding="utf-8")
    HTML_PATH.write_text(html_report(records), encoding="utf-8")
    print(f"Markdown: {MD_PATH}")
    print(f"HTML: {HTML_PATH}")


if __name__ == "__main__":
    main()
